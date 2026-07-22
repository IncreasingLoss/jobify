"""
Jobify — FastAPI Routes
"""

import io
import json
import threading
from pathlib import Path
import re
import zipfile
from typing import Optional

import numpy as np
import pandas as pd
from fastapi import FastAPI, File, Query, Request, UploadFile, Form
from fastapi.responses import HTMLResponse, JSONResponse, Response
import uvicorn

from functions import (
    calculate_max_workers,
    check_ollama,
    extract_text_from_file, 
    generate_google_query,
    generate_variants,
    load_skills_from_file,
    lock,
    parse_skills_csv_bytes,
    run_classify,
    run_match_jobs, 
    run_scrape,
    save_skills_to_file,
    state,
)

app = FastAPI(title="Jobify")
HTML_PATH = Path(__file__).parent / "front_end.html"


@app.get("/", response_class=HTMLResponse)
def index():
    if not HTML_PATH.exists():
        return HTMLResponse(
            "<h1>front_end.html not found</h1><p>Place it next to api.py</p>",
            status_code=500,
        )
    return HTML_PATH.read_text(encoding="utf-8")


# ─── Ollama ─────────────────────────────────────────────────────────────────

@app.get("/api/ollama-status")
def ollama_status():
    with lock:
        return {
            "connected": state["ollama_connected"],
            "models": state["ollama_models"],
            "error": state["ollama_error"],
        }


@app.post("/api/generate-query")
async def gen_query(req: Request):
    p = await req.json()
    kw = p.get("keywords", "").strip()
    loc = p.get("location", "").strip()
    if not kw:
        return JSONResponse({"error": "Keywords required"}, status_code=400)
    model = p.get("model", state["ollama_models"][0] if state["ollama_models"] else "qwen3:8b")
    q = generate_google_query(kw, loc, model)
    if q:
        return {"query": q}
    return JSONResponse({"error": "Failed to generate query"}, status_code=500)


# ─── AI Variant Generation ─────────────────────────────────────────────────

@app.post("/api/generate-variants")
async def gen_variants(req: Request):
    p = await req.json()
    kw = p.get("keywords", "").strip()
    loc = p.get("location", "").strip()
    if not kw:
        return JSONResponse({"error": "Keywords required"}, status_code=400)
    if not loc:
        return JSONResponse({"error": "Location required"}, status_code=400)
    lang = p.get("language", "English")
    model = p.get("model", state["ollama_models"][0] if state["ollama_models"] else "qwen3:8b")
    result = generate_variants(kw, loc, lang, model)
    return result


# ─── Auto Workers ──────────────────────────────────────────────────────────

@app.get("/api/auto-workers")
def auto_workers(model: str = Query("qwen3:8b")):
    return calculate_max_workers(model)


# ─── Skills Persistence ────────────────────────────────────────────────────

@app.get("/api/load-skills")
def load_skills():
    saved = load_skills_from_file()
    if saved:
        return {"skills": saved, "source": "saved"}
    return {"skills": None, "source": "none"}


@app.post("/api/save-skills")
async def save_skills(req: Request):
    p = await req.json()
    skills_list = p.get("skills")
    if not skills_list:
        return JSONResponse({"error": "No skills to save"}, status_code=400)
    ok = save_skills_to_file(skills_list)
    if ok:
        return {"ok": True}
    return JSONResponse({"error": "Failed to save"}, status_code=500)


@app.post("/api/parse-skills-csv")
async def parse_skills_csv(file: UploadFile = File(...)):
    try:
        content = await file.read()
        parsed = parse_skills_csv_bytes(content)
        if parsed:
            return {"skills": parsed, "count": len(parsed)}
        return JSONResponse(
            {"error": "Could not parse any skills from this file"}, status_code=400
        )
    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=500)


# ─── Scraping ───────────────────────────────────────────────────────────────

@app.post("/api/scrape")
async def start_scrape(req: Request):
    with lock:
        if state["scrape_status"] == "running":
            return JSONResponse({"error": "Scrape already running"}, status_code=409)
    p = await req.json()
    kw = p.get("keywords", "").strip()
    if not kw:
        return JSONResponse({"error": "Keywords required"}, status_code=400)
    threading.Thread(target=run_scrape, args=(p,), daemon=True).start()
    return {"ok": True}


@app.get("/api/scrape-status")
def scrape_status():
    with lock:
        return {
            "status": state["scrape_status"],
            "progress": state["scrape_progress"],
            "site_status": state["scrape_site_status"],
            "total_raw": state["scrape_total_raw"],
            "total_deduped": state["scrape_total_deduped"],
            "duplicates": state["scrape_duplicates"],
        }


@app.get("/api/scrape-data")
def scrape_data(page: int = Query(1, ge=1), per_page: int = Query(50, ge=1, le=200)):
    with lock:
        df = state["scraped_df"]
    if df is None or df.empty:
        return {"data": [], "total": 0, "page": 1, "pages": 0}
        
    strict_order = [
        "title", "company", "location", "date_posted", 
        "site", "job_url", "id", "emails", "description", "company_url", "job_type"
    ]
    existing_cols = [c for c in strict_order if c in df.columns]
    other_cols = [c for c in df.columns if c not in strict_order]
    df = df[existing_cols + other_cols]
    
    total = len(df)
    pages = max(1, (total + per_page - 1) // per_page)
    start = (page - 1) * per_page
    return {"data": _ser(df.iloc[start:start + per_page]),
            "total": total, "page": page, "pages": pages}


@app.post("/api/load-csv")
async def load_csv(file: UploadFile = File(...)):
    try:
        content = await file.read()
        df = pd.read_csv(io.BytesIO(content))
        if "title" not in df.columns:
            return JSONResponse({"error": "CSV must have a 'title' column"}, status_code=400)
        raw = len(df)
        if "description" not in df.columns:
            df["description"] = np.nan
        df["description"] = df["description"].replace(r"^\s*$", np.nan, regex=True)
        df["_tk"] = df["title"].str.strip().str.lower()
        df["_ck"] = df["company"].str.strip().str.lower() if "company" in df.columns else ""
        df = df.drop_duplicates(subset=["_tk", "_ck"], keep="first")
        df = df.drop(columns=["_tk", "_ck"])
        cols = [c for c in [
            "id", "site", "job_url", "title", "company", "location",
            "date_posted", "emails", "description", "company_url", "job_type",
        ] if c in df.columns]
        df = df.loc[:, cols]
        with lock:
            state["scraped_df"] = df
            state["scrape_total_raw"] = raw
            state["scrape_total_deduped"] = len(df)
            state["scrape_duplicates"] = raw - len(df)
            state["scrape_status"] = "done"
            state["scrape_progress"] = f"Loaded {len(df)} jobs from CSV"
        return {"ok": True, "total_raw": raw, "total_deduped": len(df), "duplicates": raw - len(df)}
    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=500)


# ─── Classification ─────────────────────────────────────────────────────────

@app.post("/api/classify")
async def start_classify(req: Request):
    with lock:
        if state["classify_status"] == "running":
            return JSONResponse({"error": "Classification already running"}, status_code=409)
        if state["scraped_df"] is None or state["scraped_df"].empty:
            return JSONResponse({"error": "No scraped data. Search or load CSV first."}, status_code=400)
    p = await req.json()
    threading.Thread(target=run_classify, args=(p,), daemon=True).start()
    return {"ok": True}


@app.get("/api/classify-status")
def classify_status():
    with lock:
        return {
            "status": state["classify_status"],
            "progress": state["classify_progress"],
            "current": state["classify_current"],
            "total": state["classify_total"],
            "target_types": state.get("classify_target_types", []),
            "used_skills": state["classify_used_skills"],
            "stats": state["classify_stats"],
        }


@app.get("/api/classify-data")
def classify_data(
    page: int = Query(1, ge=1), per_page: int = Query(50, ge=1, le=200),
    min_match: int = Query(0, ge=0, le=100),
    only_matching: bool = Query(True), 
    show_rejected: bool = Query(False), 
):
    with lock:
        df = state["classified_df"]
        used_skills = state["classify_used_skills"]
    if df is None or df.empty:
        return {"data": [], "total": 0, "page": 1, "pages": 0, "used_skills": False}
    
    filtered = df.copy()
    
    if not show_rejected:
        filtered = filtered[filtered["is_right_jobtype"] == "yes"]
        
    has_sm = "skills_matching" in filtered.columns
    if has_sm and only_matching:
        mask = filtered["skills_matching"].isna() | (filtered["skills_matching"].fillna(-1) >= min_match)
        filtered = filtered[mask]
    
    strict_order = [
        "title", "company", "skills_matching", "job_categories", "location", "date_posted",
        "site", "job_url", "is_right_jobtype",
        "is_fulltime", "is_parttime", "is_working_student", "is_internship", "is_remote",
    ]
    existing_cols = [c for c in strict_order if c in filtered.columns]
    other_cols = [c for c in filtered.columns if c not in strict_order]
    filtered = filtered[existing_cols + other_cols]
    
    total = len(filtered)
    pages = max(1, (total + per_page - 1) // per_page)
    start = (page - 1) * per_page
    rows = _ser(filtered.iloc[start:start + per_page])
    if not used_skills:
        for row in rows:
            if row.get("skills_matching") is None:
                row["skills_matching"] = "No skills — couldn't match"
    return {"data": rows, "total": total, "page": page, "pages": pages, "used_skills": used_skills}


# ─── Exports ────────────────────────────────────────────────────────────────

@app.get("/api/export/scraped")
def export_scraped():
    with lock:
        df = state["scraped_df"]
    if df is None or df.empty:
        return JSONResponse({"error": "No data"}, status_code=404)
    buf = io.StringIO()
    df.to_csv(buf, index=False)
    return Response(buf.getvalue(), media_type="text/csv",
                    headers={"Content-Disposition": "attachment; filename=jobs_scraped.csv"})


@app.get("/api/export/final")
def export_final(
    min_match: int = Query(0, ge=0, le=100),
    only_matching: bool = Query(True), 
    show_rejected: bool = Query(False),
):
    with lock:
        df = state["classified_df"]
        used_skills = state["classify_used_skills"]
    if df is None or df.empty:
        return JSONResponse({"error": "No data"}, status_code=404)
    
    filtered = df.copy()
    
    if not show_rejected:
        filtered = filtered[filtered["is_right_jobtype"] == "yes"]
        
    has_sm = "skills_matching" in filtered.columns
    if has_sm and only_matching:
        mask = filtered["skills_matching"].isna() | (filtered["skills_matching"].fillna(-1) >= min_match)
        filtered = filtered[mask]
    
    if has_sm and not filtered.empty:
        filtered = filtered.copy()
        filtered = filtered.sort_values("skills_matching", ascending=False, na_position="last")
    if not used_skills and has_sm:
        filtered = filtered.copy()
        filtered["skills_matching"] = filtered["skills_matching"].fillna("No skills — couldn't match")
    buf = io.StringIO()
    filtered.to_csv(buf, index=False)
    return Response(buf.getvalue(), media_type="text/csv",
                    headers={"Content-Disposition": "attachment; filename=jobs_final.csv"})


# ─── Application Matcher ─────────────────────────────────────────────────────

@app.get("/api/matcher-jobs")
def matcher_jobs(
    min_match: int = Query(0, ge=0, le=100),
    only_matching: bool = Query(True),
    show_rejected: bool = Query(False),
):
    with lock:
        df = state.get("classified_df")
    if df is None or df.empty:
        return {"data": [], "total": 0}

    filtered = df.copy()
    
    if not show_rejected:
        if "is_right_jobtype" in filtered.columns:
            filtered = filtered[filtered["is_right_jobtype"] == "yes"]
        else:
            return {"data": [], "total": 0}
    
    if "skills_matching" in filtered.columns and only_matching:
        mask = filtered["skills_matching"].isna() | (filtered["skills_matching"].fillna(-1) >= min_match)
        filtered = filtered[mask]

    rows = []
    for idx, row in filtered.iterrows():
        r = {"_idx": int(idx)}
        for c in ("title", "company", "skills_matching", "location", "site"):
            if c in row.index:
                v = row[c]
                if pd.isna(v):
                    r[c] = None
                elif isinstance(v, (np.integer,)):
                    r[c] = int(v)
                elif isinstance(v, (np.floating,)):
                    r[c] = float(v)
                else:
                    r[c] = str(v)
        rows.append(r)
    return {"data": rows, "total": len(rows)}


@app.post("/api/match-jobs")
async def match_jobs(
    job_ids: str = Form(...),
    cv_file: Optional[UploadFile] = File(default=None),
    cv_text: Optional[str] = Form(default=None),
    cl_file: Optional[UploadFile] = File(default=None),
    cl_text: Optional[str] = Form(default=None),
    language: str = Form("English"),
    model: str = Form("qwen3:8b"),
    skills: Optional[str] = Form(default=None),
):
    with lock:
        if state["match_status"] == "running":
            return JSONResponse({"error": "Matching already running"}, status_code=409)

    try:
        indices = [int(x) for x in json.loads(job_ids)]
    except (json.JSONDecodeError, ValueError):
        return JSONResponse({"error": "Invalid job_ids"}, status_code=400)

    if not indices:
        return JSONResponse({"error": "No jobs selected"}, status_code=400)

    cv_extracted = None
    cv_ext = ".txt"
    if cv_file and cv_file.filename:
        cv_bytes = await cv_file.read()
        cv_ext = Path(cv_file.filename).suffix.lower() or ".txt"
        cv_extracted = extract_text_from_file(cv_file.filename, cv_bytes)
        if not cv_extracted:
            return JSONResponse(
                {"error": f"Could not extract text from CV file: {cv_file.filename}"},
                status_code=400,
            )
    if cv_text and cv_text.strip():
        cv_extracted = cv_text.strip()

    cl_extracted = None
    cl_ext = ".txt"
    if cl_file and cl_file.filename:
        cl_bytes = await cl_file.read()
        cl_ext = Path(cl_file.filename).suffix.lower() or ".txt"
        cl_extracted = extract_text_from_file(cl_file.filename, cl_bytes)
        if not cl_extracted:
            return JSONResponse(
                {"error": f"Could not extract text from Cover Letter file: {cl_file.filename}"},
                status_code=400,
            )
    if cl_text and cl_text.strip():
        cl_extracted = cl_text.strip()

    if not cv_extracted and not cl_extracted:
        return JSONResponse(
            {"error": "Provide at least a CV or cover letter (file or text)"},
            status_code=400,
        )

    skills_list = None
    if skills:
        try:
            skills_list = json.loads(skills)
        except json.JSONDecodeError:
            pass

    threading.Thread(
        target=run_match_jobs,
        args=(indices, cv_extracted, cl_extracted, cv_ext, cl_ext,
              skills_list, language, model),
        daemon=True,
    ).start()

    n_tasks = len(indices) * (1 if cv_extracted else 0) + \
              len(indices) * (1 if cl_extracted else 0)
    return {"ok": True, "total_tasks": n_tasks}


@app.get("/api/match-status")
def match_status():
    with lock:
        return {
            "status": state.get("match_status", "idle"),
            "progress": state.get("match_progress", ""),
            "current": state.get("match_current", 0),
            "total": state.get("match_total", 0),
        }


@app.get("/api/match-results")
def match_results():
    with lock:
        return {
            "results": state.get("match_results", {}),
            "status": state.get("match_status", "idle"),
        }


@app.get("/api/match-download-zip/{job_id}")
def match_download_zip(job_id: str):
    with lock:
        result = state["match_results"].get(job_id)
    if not result:
        return JSONResponse({"error": "Not found"}, status_code=404)

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:

        def _add(text, ext, fname):
            if not text:
                return
            if ext == ".pdf":
                ext = ".docx"
                fname = fname.replace(".pdf", ".docx")
            if ext == ".docx":
                try:
                    from docx import Document
                    doc = Document()
                    for line in text.split("\n"):
                        doc.add_paragraph(line)
                    d = io.BytesIO()
                    doc.save(d)
                    zf.writestr(fname, d.getvalue())
                    return
                except ImportError:
                    ext = ".txt"
                    fname = fname.replace(".docx", ".txt")
            zf.writestr(fname, text.encode("utf-8"))

        _add(result.get("cv_text"), result.get("cv_ext", ".txt"),
             result.get("cv_filename", "CV.txt"))
        _add(result.get("cl_text"), result.get("cl_ext", ".txt"),
             result.get("cl_filename", "CoverLetter.txt"))

    safe = re.sub(r"[^a-zA-Z0-9_-]", "_", result.get("title", "application"))
    return Response(
        buf.getvalue(),
        media_type="application/zip",
        headers={"Content-Disposition": f"attachment; filename=application_{safe}.zip"},
    )


@app.get("/api/match-download/{job_id}/{doc_type}")
def match_download_single(job_id: str, doc_type: str):
    with lock:
        result = state["match_results"].get(job_id)
    if not result:
        return JSONResponse({"error": "Not found"}, status_code=404)

    if doc_type == "cv":
        text = result.get("cv_text")
        ext = result.get("cv_ext", ".txt")
        fname = result.get("cv_filename", "CV.txt")
    elif doc_type == "cl":
        text = result.get("cl_text")
        ext = result.get("cl_ext", ".txt")
        fname = result.get("cl_filename", "CoverLetter.txt")
    else:
        return JSONResponse({"error": "Use 'cv' or 'cl'"}, status_code=400)

    if not text:
        return JSONResponse({"error": f"No {doc_type} for this job"}, status_code=404)

    if ext == ".pdf":
        ext = ".docx"
        fname = fname.replace(".pdf", ".docx")

    if ext == ".docx":
        try:
            from docx import Document
            doc = Document()
            for line in text.split("\n"):
                doc.add_paragraph(line)
            d = io.BytesIO()
            doc.save(d)
            return Response(
                d.getvalue(),
                media_type="application/vnd.openxmlformats-officedocument"
                           ".wordprocessingml.document",
                headers={"Content-Disposition": f"attachment; filename={fname}"},
            )
        except ImportError:
            ext = ".txt"
            fname = fname.replace(".docx", ".txt")

    return Response(
        text,
        media_type="text/plain; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename={fname}"},
    )


@app.post("/api/parse-matcher-file")
async def parse_matcher_file(file: UploadFile = File(...)):
    try:
        content = await file.read()
        text = extract_text_from_file(file.filename, content)
        if text:
            return {"text": text, "filename": file.filename}
        return JSONResponse(
            {"error": f"Could not extract text from {file.filename}. "
                      f"Supported: .txt .docx .pdf .pptx .odt .rtf"},
            status_code=400,
        )
    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=500)

def _ser(df_slice):
    data = []
    for _, row in df_slice.iterrows():
        r = {}
        for c in row.index:
            v = row[c]
            if pd.isna(v):
                r[c] = None
            elif isinstance(v, (np.integer,)):
                r[c] = int(v)
            elif isinstance(v, (np.floating,)):
                r[c] = float(v)
            else:
                r[c] = str(v)
        data.append(r)
    return data



if __name__ == "__main__":
    print("Starting Jobify UI on http://localhost:8000")
    threading.Thread(target=check_ollama, daemon=True).start()
    uvicorn.run(app, host="0.0.0.0", port=8000, log_level="warning")